"""Document text extraction with location metadata for RAG indexing."""
from __future__ import annotations

import csv
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".csv"}


def extract_text(path: Path) -> list[tuple[int | None, str]]:
    """Return extracted text segments and optional page/row numbers.

    PDF pages and CSV rows are numbered from 1 and 2 respectively for CSV rows
    (to account for the header row). TXT/Markdown/DOCX are returned as one segment.
    """
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [(None, path.read_text(encoding="utf-8", errors="replace"))]
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        extracted = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        # OCR is opt-in and optional: install pytesseract + pdf2image and the
        # system Tesseract/Poppler binaries to handle image-only scanned PDFs.
        if extracted and not any(text.strip() for _, text in extracted):
            try:
                from pdf2image import convert_from_path
                import pytesseract
                images = convert_from_path(str(path), dpi=180)
                extracted = [(i + 1, pytesseract.image_to_string(image) or "") for i, image in enumerate(images)]
            except ImportError:
                # Keep the standard dependency footprint small; callers can
                # explain how to enable OCR when no text layer exists.
                pass
        return extracted
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle)
            rows: list[tuple[int, str]] = []
            for row_number, row in enumerate(reader, start=2):
                fields = [
                    f"{key}: {value}" for key, value in row.items()
                    if key and value is not None and str(value).strip()
                ]
                if fields:
                    rows.append((row_number, f"Row {row_number}: " + " | ".join(fields)))
        return rows
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        parts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [(None, "\n".join(parts))]
    return []


def validate_document_size(path: Path, max_bytes: int) -> None:
    """Reject oversized documents before parsing them."""
    if max_bytes < 1:
        raise ValueError("Maximum document size must be positive.")
    size = path.stat().st_size
    if size > max_bytes:
        raise ValueError(f"{path.name} is {size / (1024 * 1024):.1f} MB; the configured limit is {max_bytes / (1024 * 1024):.0f} MB.")
